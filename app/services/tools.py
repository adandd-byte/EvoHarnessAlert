from __future__ import annotations

import smtplib  # 邮件发送
import ssl      # TLS 上下文
import threading  # 线程锁
from email.message import EmailMessage  # 邮件消息体
from pathlib import Path  # 路径操作

from openpyxl import Workbook, load_workbook  # Excel 读写
from sqlalchemy.orm import Session  # ORM 会话

from app.core.config import Settings  # 全局配置
from app.core.enums import ToolStatus  # 工具执行状态
from app.models.entities import AlertEvent, Incident, LedgerRecord, NotificationRecord  # 实体


EXCEL_WRITE_LOCK = threading.Lock()  # Excel 写入线程锁，避免并发写坏文件


class ToolOrchestrationService:
    # 工具编排服务：把告警动作落地到台账、事件、通知
    def __init__(self, db: Session, settings: Settings):
        self.db = db  # 数据库会话
        self.settings = settings  # 全局配置

    def write_ledger(self, alert: AlertEvent) -> LedgerRecord:
        # 把告警写入 Excel 台账，并登记台账记录（同告警幂等）
        existing = self.db.query(LedgerRecord).filter(LedgerRecord.alert_id == alert.id, LedgerRecord.status == ToolStatus.SUCCESS.value).first()
        if existing is not None:  # 已成功写过的直接复用
            return existing
        path = Path(self.settings.excel_path)  # 台账文件路径
        path.parent.mkdir(parents=True, exist_ok=True)  # 确保目录存在
        with EXCEL_WRITE_LOCK:  # 加锁防止并发写冲突
            if path.exists():  # 文件已存在则追加
                workbook = load_workbook(path)
                sheet = workbook.active
            else:  # 首次则新建并写表头
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "告警台账"
                sheet.append(["告警ID", "优先级", "类型", "状态", "标题", "指纹", "创建时间"])
            sheet.append([alert.id, alert.severity, alert.alert_type, alert.status, alert.title, alert.fingerprint, alert.created_at.isoformat()])  # 追加告警行
            workbook.save(path)  # 保存文件
        record = LedgerRecord(alert_id=alert.id, file_path=str(path), status=ToolStatus.SUCCESS.value, message="告警台账已写入")  # 建台账记录
        self.db.add(record)  # 落库
        self.db.commit()
        return record

    def record_incident_upsert(self, alert: AlertEvent, incident: Incident | None) -> LedgerRecord:
        # 登记"事件已创建/更新"的台账记录
        record = LedgerRecord(
            alert_id=alert.id,  # 关联告警
            file_path="incident-db",  # 落点标识
            status=ToolStatus.SUCCESS.value,
            message=f"事件已创建或更新：incidentId={incident.id if incident else '无'}",
        )
        self.db.add(record)  # 落库
        self.db.commit()
        return record

    def send_notification(self, alert: AlertEvent, incident: Incident) -> NotificationRecord:
        # 发送 P0/P1 告警通知：支持 log / smtp 两种模式，同告警幂等
        existing = self.db.query(NotificationRecord).filter(NotificationRecord.alert_id == alert.id, NotificationRecord.status == ToolStatus.SUCCESS.value).first()
        if existing is not None:  # 已发送过的直接复用
            return existing
        recipient = self.settings.alert_email_to.strip() or "log"  # 收件人；未配置则记日志
        mode = self.settings.alert_email_delivery_mode.strip().lower()  # 投递模式
        if mode == "log":  # 日志模式只记录
            return self._save(alert, incident, recipient, ToolStatus.SUCCESS.value, f"P0/P1 告警通知已记录：alertId={alert.id}, incidentId={incident.id}, deliveryMode=log")
        if mode != "smtp":  # 非法模式
            return self._save(alert, incident, recipient, ToolStatus.FAILED.value, f"未知通知模式：{self.settings.alert_email_delivery_mode}")
        missing = self._missing_email_config()  # 检查邮件配置是否齐全
        if missing:  # 缺配置则失败
            return self._save(alert, incident, recipient, ToolStatus.FAILED.value, f"邮件未发送，缺少配置：{', '.join(missing)}")
        try:  # 尝试发送邮件
            self._send_email(alert, incident)
        except Exception as exc:  # 失败记录原因
            return self._save(alert, incident, recipient, ToolStatus.FAILED.value, f"邮件发送失败：{type(exc).__name__}: {exc}")
        return self._save(alert, incident, recipient, ToolStatus.SUCCESS.value, f"P0/P1 告警邮件已发送：alertId={alert.id}")

    def _save(self, alert: AlertEvent, incident: Incident, recipient: str, status: str, message: str) -> NotificationRecord:
        # 保存一条通知记录
        record = NotificationRecord(alert_id=alert.id, incident_id=incident.id, channel="email", recipient=recipient, status=status, message=message)
        self.db.add(record)  # 落库
        self.db.commit()
        return record

    def _missing_email_config(self) -> list[str]:
        # 检查 SMTP 关联配置，返回缺失项列表
        missing = []
        if not self.settings.smtp_host.strip():  # 缺 SMTP 主机
            missing.append("SMTP_HOST")
        if not (self.settings.alert_email_from.strip() or self.settings.smtp_username.strip()):  # 缺发件人
            missing.append("ALERT_EMAIL_FROM 或 SMTP_USERNAME")
        if not self.settings.alert_email_to.strip():  # 缺收件人
            missing.append("ALERT_EMAIL_TO")
        return missing

    def _send_email(self, alert: AlertEvent, incident: Incident) -> None:
        # 组装并发送告警邮件
        message = EmailMessage()  # 新建邮件
        message["Subject"] = f"{self.settings.alert_email_subject_prefix} alertId={alert.id}"  # 主题
        message["From"] = self.settings.alert_email_from.strip() or self.settings.smtp_username.strip()  # 发件人
        message["To"] = self.settings.alert_email_to  # 收件人
        message.set_content("\n".join([  # 邮件正文
            "EvoHarnessAlert 检测到 P0/P1 高优先级告警，请尽快处理。",  # 开头
            "",  # 空行
            f"事件ID：{incident.id}",  # 事件 ID
            f"告警ID：{alert.id}",  # 告警 ID
            f"标题：{alert.title}",  # 标题
            f"优先级：{alert.severity}",  # 级别
            f"类型：{alert.alert_type}",  # 类型
            f"状态：{alert.status}",  # 状态
            f"摘要：{incident.summary}",  # 摘要
            f"影响：{incident.impact}",  # 影响
            f"处置建议：{incident.runbook_hint}",  # 建议
        ]))
        context = ssl.create_default_context()  # 默认 SSL 上下文
        if self.settings.smtp_use_ssl:  # SSL 直连
            with smtplib.SMTP_SSL(self.settings.smtp_host, self.settings.smtp_port, timeout=self.settings.smtp_timeout_seconds, context=context) as server:
                self._send(server, message)  # 登录并发送
            return
        with smtplib.SMTP(self.settings.smtp_host, self.settings.smtp_port, timeout=self.settings.smtp_timeout_seconds) as server:  # 普通 SMTP
            server.ehlo()  # 握手
            if self.settings.smtp_use_tls:  # 启用 STARTTLS
                server.starttls(context=context)  # 升级为加密
                server.ehlo()  # 再次握手
            self._send(server, message)  # 登录并发送

    def _send(self, server: smtplib.SMTP, message: EmailMessage) -> None:
        # 有账号则登录，然后发送邮件
        if self.settings.smtp_username:  # 需要认证
            server.login(self.settings.smtp_username, self.settings.smtp_password)  # 登录
        server.send_message(message)  # 发送